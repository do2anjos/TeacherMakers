import openpyxl

def analyze_all():
    # 1. LEARNABILITY
    wb_l = openpyxl.load_workbook(r'c:\Users\Do2anjos\Desktop\T.I\TeacherMakers\Avaliação de Primeiro Acesso e Aprendizagem Inicial (Learnability) — TeacherMakers (respostas).xlsx')
    ws_l = wb_l.active
    rows_l = list(ws_l.iter_rows(values_only=True))[1:]
    print("=== LEARNABILITY ===")
    print(f"Total participantes: {len(rows_l)}")
    # Q4: finalidade correta?
    corretos_finalidade = sum(1 for r in rows_l if "Cultura Maker e Pensamento Computacional" in str(r[4]))
    print(f"Acerto da finalidade: {corretos_finalidade}/{len(rows_l)} ({corretos_finalidade/len(rows_l)*100:.1f}%)")
    # Q10: Ajuda externa necessária?
    sem_ajuda = sum(1 for r in rows_l if "sozinho" in str(r[10]))
    print(f"Navegou sem ajuda externa: {sem_ajuda}/{len(rows_l)} ({sem_ajuda/len(rows_l)*100:.1f}%)")
    # Q11: Tempo para entender
    tempos = [r[11] for r in rows_l]
    print(f"Tempos de compreensão: {tempos}")

    # 2. USABILIDADE
    wb_u = openpyxl.load_workbook(r'c:\Users\Do2anjos\Desktop\T.I\TeacherMakers\Ficha de Teste de Usabilidade — Tarefa Crítica (Módulo de Atividades) (respostas).xlsx')
    ws_u = wb_u.active
    rows_u = list(ws_u.iter_rows(values_only=True))[1:]
    print("\n=== USABILIDADE (TAREFA CRÍTICA) ===")
    print(f"Total participantes: {len(rows_u)}")
    sucessos = sum(1 for r in rows_u if "Sucesso Total sem Ajuda" in str(r[4]))
    print(f"Taxa de sucesso sem ajuda: {sucessos}/{len(rows_u)} ({sucessos/len(rows_u)*100:.1f}%)")
    tempos_u = [r[5] for r in rows_u]
    erros_u = [r[6] for r in rows_u]
    dificuldades_u = [r[8] for r in rows_u]
    print(f"Tempos de conclusão: {tempos_u}")
    print(f"Erros cometidos: {erros_u}")
    print(f"Dificuldades relatadas: {dificuldades_u}")
    notas_facilidade = [r[9] for r in rows_u]
    print(f"Notas de facilidade percebida (1-5): {notas_facilidade}")

    # 3. TÉCNICA
    wb_t = openpyxl.load_workbook(r'c:\Users\Do2anjos\Desktop\T.I\TeacherMakers\Checklist de Avaliação Técnica e Segurança da API — TeacherMakers (respostas).xlsx')
    ws_t = wb_t.active
    rows_t = list(ws_t.iter_rows(values_only=True))[1:]
    print("\n=== TÉCNICA ===")
    print(f"Total avaliações: {len(rows_t)}")
    for i, r in enumerate(rows_t, 1):
        print(f"Avaliador {i} ({r[1]}, env={r[2]}):")
        print(f"  Cenário 1 (Sem foto): {r[3][:25]}...")
        print(f"  Cenário 2 (Sem texto): {r[4][:25]}...")
        print(f"  Cenário 3 (Formato inválido): {r[5][:25]}...")
        print(f"  Cenário 4 (Token manipulado): {r[6][:25]}...")
        print(f"  Cenário 5 (Pular pré-req): {r[7][:25]}...")
        print(f"  Consistência DB: {r[8]}")
        print(f"  Parecer: {r[9][:30]}...")

    # 4. INSPEÇÃO VISUAL
    wb_i = openpyxl.load_workbook(r'c:\Users\Do2anjos\Desktop\T.I\TeacherMakers\Ficha de Inspeção Visual e Percurso Cognitivo — TeacherMakers (respostas).xlsx')
    ws_i = wb_i.active
    rows_i = list(ws_i.iter_rows(values_only=True))[1:]
    print("\n=== INSPEÇÃO VISUAL ===")
    print(f"Total avaliadores: {len(rows_i)}")
    for i, r in enumerate(rows_i, 1):
        print(f"Avaliador {i} ({r[1]}): Tela={r[2]}")
        print(f"  Q1 (Sabe o que fazer?): {r[3]}")
        print(f"  Q2 (Vê como fazer?): {r[4]}")
        print(f"  Q3 (Entende que leva ao objetivo?): {r[5]}")
        print(f"  Q4 (Recebe feedback adequado?): {r[6]}")
        print(f"  Notas: Hierarquia={r[7]}, Legibilidade={r[8]}, Consistência={r[9]}, Clareza={r[10]}")
        print(f"  Problema: {r[11]} - {r[12]} (Severidade: {r[13]})")

    # 5. PERFORMANCE
    wb_p = openpyxl.load_workbook(r'c:\Users\Do2anjos\Desktop\T.I\TeacherMakers\Registro de Auditoria de Desempenho (Performance) — TeacherMakers (respostas).xlsx')
    ws_p = wb_p.active
    rows_p = list(ws_p.iter_rows(values_only=True))[1:]
    print("\n=== PERFORMANCE ===")
    r_p = rows_p[0]
    print(f"Avaliador: {r_p[1]}, URL/Tela: {r_p[2]}, Perfil: {r_p[3]}")
    print(f"Score Geral: {r_p[5]}/100")
    print(f"LCP: {r_p[6]}s, FCP: {r_p[7]}s, TBT: {r_p[8]}ms, CLS: {r_p[9]}")
    print(f"Alertas/Gargalos: {r_p[10]}")
    print(f"Parecer: {r_p[11]}")

    # 6. EDUCACIONAL
    wb_e = openpyxl.load_workbook(r'c:\Users\Do2anjos\Desktop\T.I\TeacherMakers\Avaliação de Conhecimento em Cultura Maker e Pensamento Computacional — Pré-teste (respostas).xlsx')
    ws_e = wb_e.active
    rows_e = list(ws_e.iter_rows(values_only=True))[1:]
    print("\n=== EDUCACIONAL ===")
    print(f"Total respostas: {len(rows_e)}")
    r_e = rows_e[0]
    print(f"Participante: {r_e[2]}, Momento: {r_e[3]}, Nota: {r_e[1]}/10")
    # check questions
    for q_idx in range(4, 14):
        print(f"  Questão {q_idx-3}: {ws_e.cell(row=1, column=q_idx+1).value.splitlines()[0]}")
        print(f"    Resposta: {r_e[q_idx][:60]}...")

if __name__ == '__main__':
    analyze_all()
