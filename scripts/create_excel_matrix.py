import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_excel_matrix():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Problemas Consolidados'

    # Title
    ws.merge_cells('A1:I1')
    ws['A1'] = 'MATRIZ CONSOLIDADA DE PROBLEMAS — TEACHERMAKERS (v1.0)'
    ws['A1'].font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 35

    headers = [
        'ID', 'Dimensão de Origem', 'Tela / Componente', 'Descrição do Problema',
        'Severidade', 'Causa Raiz', 'Impacto Operacional / Pedagógico', 'Ação Recomendada (Solução)', 'Prioridade'
    ]

    ws.append([]) # row 2 empty
    ws.append(headers) # row 3 headers
    ws.row_dimensions[3].height = 28

    header_fill = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')

    for col_idx in range(1, 10):
        cell = ws.cell(row=3, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    data = [
        [
            'PB-01',
            'Avaliação Técnica (API)',
            'Endpoint /progress/pratica/finish',
            'Gravação de registros parciais no banco de dados em requisições anômalas após falha na validação multipart.',
            'Crítica',
            'Ausência de transação atômica (BEGIN/COMMIT/ROLLBACK) no backend.',
            'Corrupção da base de dados, risco de avanço indevido e emissão de certificados sem evidência real.',
            'Implementar transação estrita no SQLite/ORM com ROLLBACK em qualquer falha de validação.',
            'P1 (Imediata)'
        ],
        [
            'PB-02',
            'Acessibilidade (Lighthouse)',
            'Aplicação Global',
            'Pontuação de acessibilidade crítica (Score 50/100). Contraste insuficiente e falta de suporte para teclado.',
            'Alta',
            'Não conformidade com diretrizes da WCAG 2.1 AA na folha de estilos.',
            'Exclusão de educadores com deficiência visual ou motora que dependem de leitores de tela ou teclado.',
            'Adequar contraste para proporção 4.5:1, adicionar aria-label e anéis de foco (:focus-visible).',
            'P2 (Curto Prazo)'
        ],
        [
            'PB-03',
            'Inspeção Visual / Usabilidade',
            'Viewport Tablet (768px a 1024px)',
            'Incoerência e quebra de layout em tablets. Cards desalinhados e digitação desconfortável do diário.',
            'Alta',
            'Falta de media queries específicas para a faixa intermediária de tablets.',
            'Tempo de tarefa dilatado (30 min), 2 erros operacionais e frustração do usuário.',
            'Implementar layout responsivo específico para tablets com campos de texto 100% width e botões >= 48px.',
            'P2 (Curto Prazo)'
        ],
        [
            'PB-04',
            'Inspeção Visual / Usabilidade',
            'Dashboard Principal',
            'Competição de estímulos visuais e hesitação na localização do Módulo 1 (início da trilha).',
            'Média',
            'Cards de módulos com o mesmo peso visual e ausência de destaque hierárquico no módulo ativo.',
            'Atrito cognitivo inicial e cliques fora do caminho para descobrir onde começar.',
            'Destacar o Módulo 1 com badge "Comece por Aqui" e borda ativa; esmaecer módulos bloqueados.',
            'P3 (Médio Prazo)'
        ],
        [
            'PB-05',
            'Usabilidade / Inspeção Visual',
            'Módulo de Atividades',
            'Feedback de confirmação de envio sutil e passageiro (toast desaparece rápido demais).',
            'Média',
            'Timeout curto da notificação sem retenção de foco visual.',
            'Sensação de insegurança do cursista quanto à gravação da atividade prática.',
            'Substituir o toast por Modal Centralizado de Sucesso com botão de avanço e atualização de progresso.',
            'P3 (Médio Prazo)'
        ],
        [
            'PB-06',
            'Teste de Usabilidade',
            'Fluxo de Cadastro / Onboarding',
            'Módulo introdutório marcado como concluído no primeiro login de novo usuário.',
            'Média',
            'Flag de status inicial incorreta atribuída no seed/migration do banco.',
            'Quebra da sequência pedagógica e estranheza relatada pelo participante.',
            'Corrigir lógica de inicialização de novo usuário para definir todos os módulos como locked/available.',
            'P3 (Médio Prazo)'
        ],
        [
            'PB-07',
            'Experiência do Usuário (UEQ-S)',
            'Interface Global',
            'Qualidade Hedônica neutra (+0,31) e percepção de sistema convencional e tradicional.',
            'Baixa',
            'Estética sóbria focada estritamente na função, sem identidade visual maker.',
            'Menor engajamento emocional e empolgação a longo prazo.',
            'Adicionar microinterações, gamificação leve (conquistas maker) e paleta visual vibrante.',
            'P4 (Evolutivo)'
        ],
        [
            'PB-08',
            'Avaliação Educacional',
            'Instrumento de Aprendizagem',
            'Ausência de dados de pré-teste pareado (apenas 1 resposta de pós-teste coletada).',
            'Baixa (Metod.)',
            'Falha de aplicação do protocolo de campo antes do início das aulas.',
            'Impossibilita cálculo do Teste T de Student para comprovação inferencial de ganho.',
            'Integrar questionário diagnóstico obrigatoriamente dentro da plataforma antes do Módulo 1.',
            'P4 (Metodológico)'
        ]
    ]

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    sev_fills = {
        'Crítica': PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid'),
        'Alta': PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid'),
        'Média': PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid'),
        'Baixa': PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid'),
        'Baixa (Metod.)': PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
    }

    prio_fills = {
        'P1 (Imediata)': PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid'),
        'P2 (Curto Prazo)': PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid'),
        'P3 (Médio Prazo)': PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid'),
        'P4 (Evolutivo)': PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid'),
        'P4 (Metodológico)': PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
    }

    for row_idx, r_data in enumerate(data, 4):
        ws.append(r_data)
        ws.row_dimensions[row_idx].height = 42
        for col_idx in range(1, 10):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name='Calibri', size=10)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if col_idx in [1, 5, 9]:
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            if col_idx == 5 and r_data[4] in sev_fills:
                cell.fill = sev_fills[r_data[4]]
                cell.font = Font(name='Calibri', size=10, bold=True)
            if col_idx == 9 and r_data[8] in prio_fills:
                cell.fill = prio_fills[r_data[8]]
                cell.font = Font(name='Calibri', size=10, bold=True)

    col_widths = {1: 12, 2: 24, 3: 26, 4: 44, 5: 16, 6: 36, 7: 38, 8: 44, 9: 18}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    out_path = r'c:\Users\Do2anjos\Desktop\T.I\TeacherMakers\Matriz_Consolidada_de_Problemas.xlsx'
    wb.save(out_path)
    print('Planilha gerada com sucesso:', out_path)

if __name__ == '__main__':
    build_excel_matrix()
